"""Turn an uploaded company document into plain text the council can use as evidence.

Rules (docs/flow.md): read in memory, never written to disk; text is DATA, never instructions;
CSV is read by code (numbers stay exact); limits keep one upload from eating the token budget.
"""

import csv
import re
import io
import uuid

from docx import Document
import openpyxl
from pypdf import PdfReader

MAX_FILES = 3
MAX_BYTES = 5 * 1024 * 1024
MAX_PDF_PAGES = 10
MAX_CHARS = 20_000  # per document, before it is chunked into evidence
CSV_ROWS = 200

KINDS = {".pdf": "pdf", ".csv": "csv", ".txt": "text", ".md": "text", ".docx": "docx", ".xlsx": "xlsx"}


class DocumentError(ValueError):
    pass


def _kind(name: str) -> str:
    ext = name[name.rfind("."):].lower() if "." in name else ""
    if ext not in KINDS:
        raise DocumentError(f"{name}: unsupported file type (use PDF, CSV, XLSX, TXT, MD or DOCX)")
    return KINDS[ext]


def _pdf(data: bytes) -> tuple[str, int]:
    reader = PdfReader(io.BytesIO(data))
    pages = reader.pages[:MAX_PDF_PAGES]
    text = "\n".join(f"[page {i + 1}] {p.extract_text() or ''}" for i, p in enumerate(pages))
    return text, len(reader.pages)


def _csv(data: bytes) -> str:
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))))
    if not rows:
        return ""
    header, body = rows[0], rows[1 : CSV_ROWS + 1]
    lines = [" | ".join(header)] + [" | ".join(r) for r in body]
    more = f"\n({len(rows) - 1 - len(body)} more rows not shown)" if len(rows) - 1 > len(body) else ""
    return f"Table with columns: {', '.join(header)}\n" + "\n".join(lines) + more


def _number(v) -> float | None:
    """Spreadsheet number rules: real numbers as-is; text like '₹1,200', 'Rs 5 Cr', '12%', '(38.4)'
    (accounting negative) or '−3' (unicode minus) parsed; anything else is not a number."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("\u20b9", "").replace("\u2212", "-")
    s = re.sub(r"(?i)^rs\.?|\b(cr|crore|lakh|lakhs|l|k|m|mn|bn)\b|%", "", s).strip()
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").strip()
    try:
        n = float(s)
    except ValueError:
        return None
    return -n if neg else n


def table_and_stats(text: str) -> tuple[dict, dict]:
    """Exact numbers by code (flow.md): per numeric column first/last/min/max/change/trend."""
    return rows_table_and_stats(list(csv.reader(io.StringIO(text))))


def clean_rows(rows: list[list]) -> list[list]:
    """Spreadsheet rules: drop fully empty rows and columns, start at the header row
    (first row with 2+ filled cells, mostly text), turn dates into ISO strings."""
    def cell(v):
        if v is None:
            return ""
        if hasattr(v, "isoformat"):
            return v.isoformat()[:10]
        return v.strip() if isinstance(v, str) else v
    rows = [[cell(c) for c in r] for r in rows]
    rows = [r for r in rows if any(c != "" for c in r)]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    keep = [i for i in range(width) if any(r[i] != "" for r in rows)]
    rows = [[r[i] for i in keep] for r in rows]
    for start, r in enumerate(rows):
        filled = [c for c in r if c != ""]
        if len(filled) >= 2 and sum(isinstance(c, str) and _number(c) is None for c in filled) >= len(filled) / 2:
            return rows[start:]
    return rows


def rows_table_and_stats(rows: list[list]) -> tuple[dict, dict]:
    if len(rows) < 2:
        return {"columns": rows[0] if rows else [], "rows": []}, {}
    header, body = [str(h) for h in rows[0]], rows[1 : CSV_ROWS + 1]
    stats = {}
    for i, col in enumerate(header):
        vals = [_number(r[i]) for r in body if i < len(r)]
        nums = [v for v in vals if v is not None]
        if len(nums) < max(2, 0.6 * len(body)):
            continue
        first, last = nums[0], nums[-1]
        change = round((last - first) / abs(first) * 100, 1) if first else None
        trend = "flat" if change is None or abs(change) < 2 else ("up" if change > 0 else "down")
        # recent = the last 3 rows: catches a recovery that the first-vs-last trend hides
        base = nums[-3] if len(nums) >= 3 else first
        recent = round((last - base) / abs(base) * 100, 1) if base else None
        recent_trend = "flat" if recent is None or abs(recent) < 2 else ("up" if recent > 0 else "down")
        stats[col] = {"first": first, "last": last, "min": min(nums), "max": max(nums), "change_pct": change, "trend": trend,
                      "recent_change_pct": recent, "recent_trend": recent_trend}
    # rows can be ragged (e.g. an unquoted "Rs 1,000" splits into two cells): keep it to the header's width
    typed = [[(_number(c) if header[i] in stats and _number(c) is not None else c) for i, c in enumerate(r[: len(header)])]
             for r in body]
    return {"columns": header, "rows": typed}, stats


def read_workbook(data: bytes) -> tuple[str, list[dict]]:
    """Every tab of an Excel workbook. Numbers use the saved (calculated) values, not formulas.
    Returns (text for the council, tabs with table + stats)."""
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    tabs, parts = [], []
    for ws in wb.worksheets[:10]:
        rows = clean_rows([list(r) for r in ws.iter_rows(max_row=CSV_ROWS + 50, values_only=True)])
        if not rows:
            continue
        table, stats = rows_table_and_stats(rows)
        tabs.append({"name": ws.title, "table": table, "stats": stats})
        cols = table["columns"]
        if len(cols) == 2 and not stats:  # a "Field | Value" tab: facts, not a table
            lines = [f"{r[0]}: {r[1]}" for r in table["rows"] if len(r) > 1]
        else:
            lines = [" | ".join(map(str, cols))] + [" | ".join(map(str, r)) for r in table["rows"]]
        parts.append(f"TAB '{ws.title}':\n" + _stats_brief(stats) + "\n".join(lines))
    if not tabs:
        raise DocumentError("the workbook is empty")
    return "\n\n".join(parts), tabs


def _docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        parts += [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
    return "\n".join(parts)


def extract(name: str, data: bytes) -> dict:
    if len(data) > MAX_BYTES:
        raise DocumentError(f"{name}: larger than 5 MB")
    kind = _kind(name)
    pages = None
    try:
        if kind == "pdf":
            text, pages = _pdf(data)
        elif kind == "csv":
            text = _csv(data)
        elif kind == "docx":
            text = _docx(data)
        elif kind == "xlsx":
            text, tabs = read_workbook(data)
        else:
            text = data.decode("utf-8", errors="replace")
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError(f"{name}: could not be read ({type(exc).__name__})") from exc
    text = " ".join(text.split()) if kind not in ("csv", "xlsx") else text.strip()
    if not text:
        raise DocumentError(f"{name}: no readable text found (scanned PDF?)")
    text = text[:MAX_CHARS]
    doc = {"id": uuid.uuid4().hex[:12], "name": name, "kind": kind, "chars": len(text), "pages": pages,
           "preview": text[:280], "text": text}
    if kind == "csv":
        doc["table"], doc["stats"] = table_and_stats(data.decode("utf-8-sig", errors="replace"))
        doc["text"] = _stats_brief(doc["stats"]) + doc["text"]
    elif kind == "xlsx":
        # the tab with the most numeric columns drives the charts and the v1->v2 diff
        main = max(tabs, key=lambda t: len(t["stats"]))
        doc.update(table=main["table"], stats=main["stats"], main_tab=main["name"],
                   tabs=[{"name": t["name"], "columns": t["table"]["columns"], "rows": len(t["table"]["rows"]),
                          "numeric_columns": list(t["stats"])} for t in tabs])
    return doc


def _stats_brief(stats: dict) -> str:
    """Code-computed facts put first, so the council quotes exact numbers instead of eyeballing a table."""
    if not stats:
        return ""
    lines = [f"{c}: first {s['first']:g}, latest {s['last']:g}, min {s['min']:g}, max {s['max']:g}"
             + (f", change {s['change_pct']:+g}% ({s['trend']})" if s["change_pct"] is not None else "")
             + (f", last 3 rows {s['recent_change_pct']:+g}% ({s['recent_trend']})" if s.get("recent_change_pct") is not None else "")
             for c, s in stats.items()]
    return "COMPUTED BY CODE (exact): " + "; ".join(lines) + "\n"


if __name__ == "__main__":
    d = extract("sales.csv", b"month,revenue\nJan,100\nFeb,80\n")
    assert d["kind"] == "csv" and "Jan | 100" in d["text"], d
    t = extract("notes.md", b"# Our shop\nWe sell phones in Pune.")
    assert t["kind"] == "text" and "phones" in t["text"]
    for bad, msg in ((("x.exe", b"1"), "unsupported"), (("big.txt", b"a" * (MAX_BYTES + 1)), "5 MB")):
        try:
            extract(*bad)
            raise AssertionError("should have failed")
        except DocumentError as exc:
            assert msg in str(exc)
    tb, st = table_and_stats('month,revenue\nJan,"\u20b91,000"\nFeb,800\nMar,500\n')
    assert table_and_stats("month,revenue\nJan,1,000\nFeb,800\n")[0]["rows"][0] == ["Jan", 1.0]  # ragged row: no crash
    assert st["revenue"]["first"] == 1000 and st["revenue"]["trend"] == "down" and st["revenue"]["change_pct"] == -50
    assert tb["rows"][0] == ["Jan", 1000.0]
    print("documents self-check ok")
