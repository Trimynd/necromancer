"""Spreadsheets (XLSX/CSV) read by code, never by the model: numbers stay exact (docs/flow.md).

Two uses:
- evidence text for the court (documents.py turns a workbook into a readable table);
- comparing 2-6 workbooks: a cell-by-cell diff, plus a verdict computed in code for each
  workbook, so the same numbers always give the same verdict and a changed number shows
  exactly how much it moved the verdict.
"""

import csv
import io
import re
from datetime import date, datetime

from openpyxl import load_workbook

MAX_ROWS = 500
MAX_COLS = 50
TEXT_ROWS = 200  # rows per sheet shown to the court as evidence


class SpreadsheetError(ValueError):
    pass


# ---------- reading ----------

def _cell(v):
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _trim(rows: list[list]) -> list[list]:
    rows = [r[:MAX_COLS] for r in rows if any(c is not None for c in r)][:MAX_ROWS]
    width = max((max((i + 1 for i, c in enumerate(r) if c is not None), default=0) for r in rows), default=0)
    return [r[:width] + [None] * (width - len(r)) for r in rows]


def read_sheets(name: str, data: bytes) -> list[dict]:
    """[{"sheet": name, "rows": [[cell, ...], ...]}] with empty rows dropped and widths evened."""
    lower = name.lower()
    if lower.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        rows = [[_cell(c) for c in r] for r in csv.reader(io.StringIO(text))]
        rows = [[parse_number(c) if parse_number(c) is not None else c for c in r] for r in rows]
        sheets = [{"sheet": "Sheet1", "rows": _trim(rows)}]
    elif lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise SpreadsheetError(f"{name}: not a readable Excel file ({type(exc).__name__})") from exc
        sheets = []
        for ws in wb.worksheets:
            rows = [[_cell(c) for c in row] for _, row in zip(range(MAX_ROWS * 2), ws.iter_rows(values_only=True))]
            sheets.append({"sheet": ws.title, "rows": _trim(rows)})
        wb.close()
    else:
        raise SpreadsheetError(f"{name}: use an Excel (.xlsx) or CSV file")
    sheets = [s for s in sheets if s["rows"]]
    if not sheets:
        raise SpreadsheetError(f"{name}: the spreadsheet is empty")
    return sheets


def sheets_text(sheets: list[dict]) -> str:
    """A workbook as plain tables, for the court's evidence."""
    out = []
    for s in sheets:
        rows = s["rows"]
        body = [" | ".join("" if c is None else str(c) for c in r) for r in rows[:TEXT_ROWS]]
        more = f"\n({len(rows) - TEXT_ROWS} more rows not shown)" if len(rows) > TEXT_ROWS else ""
        out.append(f"Sheet '{s['sheet']}':\n" + "\n".join(body) + more)
    return "\n\n".join(out)


_NUM = re.compile(r"^\(?-?[\d,]*\.?\d+\)?$")


def parse_number(v) -> float | None:
    """1234, '1,234', '₹1,234 Cr', '(120)' (negative), '12.5%' (-> 12.5). Anything else -> None."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"(?i)[₹$€£,\s]|rs\.?|inr|usd|crores?|cr\b|lakhs?|mn\b|%", "", str(v))
    if not s or not _NUM.match(s):
        return None
    neg = s.startswith("(") and s.endswith(")")
    n = float(s.strip("()"))
    return -abs(n) if neg else n


# ---------- tables ----------

def _header_index(rows: list[list]) -> int:
    """First row that looks like a header: 2+ text cells."""
    for i, r in enumerate(rows[:10]):
        if sum(1 for c in r if isinstance(c, str) and parse_number(c) is None) >= 2:
            return i
    return 0


def table(rows: list[list]) -> tuple[list[str], list[list]]:
    h = _header_index(rows)
    header = [str(c) if c is not None else f"Column {i + 1}" for i, c in enumerate(rows[h])] if rows else []
    return header, rows[h + 1:]


def _label(row: list, i: int) -> str:
    return str(row[0]) if row and row[0] is not None else f"row {i + 1}"


# ---------- diff ----------

def _pair_sheets(a: list[dict], b: list[dict]) -> list[tuple[dict, dict]]:
    by_name = {s["sheet"].lower(): s for s in b}
    pairs = [(s, by_name[s["sheet"].lower()]) for s in a if s["sheet"].lower() in by_name]
    if not pairs and a and b:  # different sheet names: compare first sheet with first sheet
        pairs = [(a[0], b[0])]
    return pairs


def diff_sheets(a: list[dict], b: list[dict]) -> dict:
    """Cell-level changes between two workbooks. Rows pair by their first-column label when the
    labels mostly match (same periods/items), otherwise by position (Q1 vs Q1 of another year)."""
    changes, added_rows, removed_rows, added_cols, removed_cols, paired_by = [], [], [], [], [], "label"
    for sa, sb in _pair_sheets(a, b):
        ha, ra = table(sa["rows"])
        hb, rb = table(sb["rows"])
        cols_b = {c.lower(): j for j, c in enumerate(hb)}
        shared = [(c, j, cols_b[c.lower()]) for j, c in enumerate(ha) if c.lower() in cols_b]
        added_cols += [f"{sb['sheet']}: {c}" for c in hb if c.lower() not in {x.lower() for x in ha}]
        removed_cols += [f"{sa['sheet']}: {c}" for c in ha if c.lower() not in cols_b]

        la = [_label(r, i) for i, r in enumerate(ra)]
        lb = [_label(r, i) for i, r in enumerate(rb)]
        overlap = len(set(la) & set(lb))
        if overlap and overlap >= min(len(la), len(lb)) / 2:
            idx_b = {l: i for i, l in enumerate(lb)}
            pairs = [(i, idx_b[l]) for i, l in enumerate(la) if l in idx_b]
            added_rows += [f"{sb['sheet']}: {l}" for l in lb if l not in set(la)]
            removed_rows += [f"{sa['sheet']}: {l}" for l in la if l not in idx_b]
        else:
            paired_by = "position"
            pairs = list(zip(range(len(ra)), range(len(rb))))
            added_rows += [f"{sb['sheet']}: {l}" for l in lb[len(ra):]]
            removed_rows += [f"{sa['sheet']}: {l}" for l in la[len(rb):]]

        for i, k in pairs:
            row_a, row_b = ra[i], rb[k]
            for col, ja, jb in shared:
                if ja == 0 and paired_by == "position" and row_a[0] != row_b[0]:
                    continue  # period labels differ by design when pairing by position
                va = row_a[ja] if ja < len(row_a) else None
                vb = row_b[jb] if jb < len(row_b) else None
                na, nb = parse_number(va), parse_number(vb)
                if na is not None and nb is not None:
                    if abs(na - nb) < 1e-9:
                        continue
                    changes.append({"sheet": sa["sheet"], "row": la[i] if la[i] == lb[k] else f"{la[i]} → {lb[k]}",
                                    "column": col, "a": na, "b": nb, "delta": round(nb - na, 4),
                                    "pct": round((nb - na) / abs(na) * 100, 1) if na else None})
                elif va != vb:
                    changes.append({"sheet": sa["sheet"], "row": la[i] if la[i] == lb[k] else f"{la[i]} → {lb[k]}",
                                    "column": col, "a": va, "b": vb, "delta": None, "pct": None})
    return {"changes": changes, "total_changed": len(changes), "paired_by": paired_by,
            "added_rows": added_rows, "removed_rows": removed_rows,
            "added_columns": added_cols, "removed_columns": removed_cols}


# ---------- financials ----------

# Order matters: the first pattern a header matches wins ("customer acquisition cost" is CAC, not cost).
METRICS = [
    ("cac", r"\bcac\b|acquisition cost"),
    ("ebitda", r"ebitda"),
    ("net_profit", r"net\s*(profit|income|loss|earnings)|profit|\bpat\b|earnings|loss"),
    ("revenue", r"revenue|sales|turnover|income from operations|\bgmv\b"),
    ("expenses", r"expense|cost|opex|spend|burn"),
    ("customers", r"customer|user|subscriber|\bmtu\b|client|buyer|order"),
    ("cash", r"\bcash\b|bank balance"),
    ("employees", r"employee|headcount|staff"),
]
METRIC_LABELS = {"revenue": "Revenue", "expenses": "Expenses", "net_profit": "Net profit", "ebitda": "EBITDA",
                 "customers": "Customers", "cash": "Cash", "employees": "Employees", "cac": "CAC"}


def _metric(name: str) -> str | None:
    low = name.lower()
    return next((key for key, pat in METRICS if re.search(pat, low)), None)


def _series_from(header: list[str], body: list[list]) -> tuple[list[str], dict[str, list]]:
    """Columns are metrics, rows are periods."""
    cols = {}
    for j, h in enumerate(header[1:], start=1):
        m = _metric(h)
        if m and m not in cols:
            cols[m] = j
    periods = [_label(r, i) for i, r in enumerate(body)]
    series = {m: [parse_number(r[j]) if j < len(r) else None for r in body] for m, j in cols.items()}
    return periods, series


def financials(sheets: list[dict]) -> dict:
    """Find the period-by-metric table in a workbook. Handles both layouts: metrics as columns
    (periods down the side) or metrics as rows (periods across the top)."""
    best = {"sheet": None, "periods": [], "series": {}}
    for s in sheets:
        header, body = table(s["rows"])
        candidates = [_series_from(header, body)]
        # transposed: first column holds metric names, header holds periods
        if body:
            t_header = [header[0]] + [_label(r, i) for i, r in enumerate(body)]
            t_body = [[header[j]] + [r[j] if j < len(r) else None for r in body] for j in range(1, len(header))]
            candidates.append(_series_from(t_header, t_body))
        for periods, series in candidates:
            series = {m: v for m, v in series.items() if sum(x is not None for x in v) >= 1}
            if len(series) > len(best["series"]) or (series and len(series) == len(best["series"])
                                                    and len(periods) > len(best["periods"])):
                best = {"sheet": s["sheet"], "periods": periods, "series": series}
    return best


# ---------- verdict (computed in code, never by the model) ----------

def _clamp(x: float) -> float:
    return round(max(0.0, min(10.0, x)), 1)


def _vals(series: dict, key: str) -> list[float]:
    return [v for v in series.get(key, []) if v is not None]


def _growth(values: list[float]) -> float | None:
    """Average growth per period from first to last value."""
    if len(values) < 2 or values[0] <= 0 or values[-1] <= 0:
        if len(values) >= 2 and values[0]:
            return (values[-1] - values[0]) / abs(values[0]) / (len(values) - 1)
        return None
    return (values[-1] / values[0]) ** (1 / (len(values) - 1)) - 1


VERDICTS = [  # (min score, label, mode, headline)
    (70, "SCALE", "Grow", "The numbers back expansion: growth and margins are both working."),
    (55, "GROW WITH CARE", "Grow", "A working business with one weak spot to fix before pushing harder."),
    (40, "RECOVERABLE", "Rescue", "Real revenue, but the economics are leaking. Fix the leak before spending to grow."),
    (0, "CRITICAL", "Rescue", "The numbers are moving the wrong way. Cut losses and find what still works."),
]

DIMENSION_LABELS = {"growth": "Revenue growth", "profitability": "Profitability", "momentum": "Margin momentum",
                    "efficiency": "Cost efficiency", "traction": "Customer traction", "stability": "Consistency",
                    "yoy": "Vs previous workbook"}


def verdict(fin: dict, prior: dict | None = None) -> dict:
    """prior: the previous workbook's financials (e.g. last year). Used only when its periods differ
    from this one's, to judge year-over-year change, which an in-year trend can't see."""
    s = fin["series"]
    rev, exp, prof, cust = _vals(s, "revenue"), _vals(s, "expenses"), _vals(s, "net_profit"), _vals(s, "customers")
    dims: dict[str, float] = {}
    kpis: dict[str, float | None] = {}

    if rev:
        kpis["total_revenue"] = round(sum(rev), 2)
        kpis["latest_revenue"] = rev[-1]
    g = _growth(rev)
    if g is not None:
        kpis["revenue_growth_pct"] = round(g * 100, 1)
        dims["growth"] = _clamp(4 + g * 40)  # 0%/period -> 4, +15%/period -> 10, -10%/period -> 0
        ups = sum(1 for x, y in zip(rev, rev[1:]) if y >= x)
        dims["stability"] = _clamp(10 * ups / (len(rev) - 1))

    # Margin per period: net profit if given, else revenue - expenses.
    pairs = list(zip(s.get("revenue", []), s.get("net_profit", []))) if prof else \
        [(r, (r - e) if r is not None and e is not None else None)
         for r, e in zip(s.get("revenue", []), s.get("expenses", []))] if exp else []
    margins = [p / r for r, p in pairs if r and p is not None]
    if margins:
        total_r = sum(r for r, p in pairs if r and p is not None)
        total_p = sum(p for r, p in pairs if r and p is not None)
        m = total_p / total_r
        kpis["net_margin_pct"] = round(m * 100, 1)
        kpis["total_profit"] = round(total_p, 2)
        kpis["latest_profit"] = next(p for r, p in reversed(pairs) if r and p is not None)
        dims["profitability"] = _clamp(5 + m * 20 if m < 0 else 5 + m * 25)  # -25% -> 0, 0 -> 5, +20% -> 10
        if len(margins) >= 2:
            kpis["margin_change_pp"] = round((margins[-1] - margins[0]) * 100, 1)
            dims["momentum"] = _clamp(5 + (margins[-1] - margins[0]) * 25)  # +20pp -> 10, -20pp -> 0
    if exp and rev:
        ratio = sum(exp) / sum(rev[:len(exp)] or rev)
        kpis["expense_ratio"] = round(ratio, 2)
        dims["efficiency"] = _clamp(10 - (ratio - 0.7) * 12)  # 0.7 -> 10, 1.0 -> 6.4, 1.5 -> 0.4
    prior_rev = _vals(prior["series"], "revenue") if prior else []
    if rev and prior_rev and sum(prior_rev) > 0 and not set(fin["periods"]) & set(prior["periods"]):
        yoy = sum(rev) / sum(prior_rev) - 1
        kpis["yoy_revenue_pct"] = round(yoy * 100, 1)
        dims["yoy"] = _clamp(5 + yoy * 12.5)  # flat -> 5, +40% -> 10, -40% -> 0
    gc = _growth(cust)
    if gc is not None:
        kpis["latest_customers"] = cust[-1]
        kpis["customer_growth_pct"] = round(gc * 100, 1)
        dims["traction"] = _clamp(4 + gc * 40)

    overall = round(sum(dims.values()) / len(dims) * 10) if dims else None
    label, mode, headline = next(((l, m, h) for floor, l, m, h in VERDICTS if overall is not None and overall >= floor),
                                 ("NOT ENOUGH DATA", None, "Add a revenue column (and ideally profit or expenses) "
                                                         "with one row per period."))
    return {"overall": overall, "label": label, "mode": mode, "headline": headline, "dimensions": dims, "kpis": kpis}


KPI_LABELS = {"total_revenue": "Total revenue", "latest_revenue": "Latest-period revenue",
              "revenue_growth_pct": "Revenue growth / period (%)", "net_margin_pct": "Net margin (%)",
              "total_profit": "Total net profit", "latest_profit": "Latest-period profit",
              "margin_change_pp": "Margin change, first → last (pp)", "expense_ratio": "Expenses ÷ revenue",
              "latest_customers": "Latest customers", "customer_growth_pct": "Customer growth / period (%)",
              "yoy_revenue_pct": "Revenue vs previous workbook (%)"}

_WHY = {"growth": "revenue_growth_pct", "profitability": "net_margin_pct", "momentum": "margin_change_pp",
        "efficiency": "expense_ratio", "traction": "customer_growth_pct",
        "yoy": "yoy_revenue_pct"}


def _fmt(v, digits: int = 1) -> str:
    return "—" if v is None else f"{v:,.{digits}f}" if isinstance(v, float) else f"{v:,}"


def verdict_change(va: dict, vb: dict) -> dict:
    keys = list(dict.fromkeys([*va["dimensions"], *vb["dimensions"]]))
    dims = [{"key": k, "label": DIMENSION_LABELS.get(k, k), "a": va["dimensions"].get(k), "b": vb["dimensions"].get(k),
             "delta": round(vb["dimensions"][k] - va["dimensions"][k], 1)
             if k in va["dimensions"] and k in vb["dimensions"] else None} for k in keys]
    drivers = []
    for d in sorted((d for d in dims if d["delta"]), key=lambda d: -abs(d["delta"])):
        kpi = _WHY.get(d["key"])
        digits = 2 if kpi == "expense_ratio" else 1
        why = (f": {KPI_LABELS[kpi].lower()} went {_fmt(va['kpis'].get(kpi), digits)} → "
               f"{_fmt(vb['kpis'].get(kpi), digits)}" if kpi else "")
        drivers.append(f"{d['label']} {'rose' if d['delta'] > 0 else 'fell'} {d['a']} → {d['b']} "
                       f"({d['delta']:+.1f}/10){why}.")
    for d in dims:
        if d["b"] is not None and d["a"] is None:
            kpi = _WHY.get(d["key"])
            why = f" ({KPI_LABELS[kpi].lower()}: {_fmt(vb['kpis'].get(kpi))})" if kpi else ""
            drivers.append(f"{d['label']} scored {d['b']}/10{why}; only the later workbook has it.")
    kpis = [{"key": k, "label": KPI_LABELS[k], "a": va["kpis"].get(k), "b": vb["kpis"].get(k),
             "delta": round(vb["kpis"][k] - va["kpis"][k], 2) if k in va["kpis"] and k in vb["kpis"] else None,
             "pct": round((vb["kpis"][k] - va["kpis"][k]) / abs(va["kpis"][k]) * 100, 1)
             if va["kpis"].get(k) and k in vb["kpis"] else None}
            for k in KPI_LABELS if k in va["kpis"] or k in vb["kpis"]]
    oa, ob = va["overall"], vb["overall"]
    return {"overall": {"a": oa, "b": ob, "delta": ob - oa if oa is not None and ob is not None else None},
            "label": {"a": va["label"], "b": vb["label"], "changed": va["label"] != vb["label"]},
            "dimensions": dims, "kpis": kpis, "drivers": drivers}


def analyze(name: str, data: bytes) -> dict:
    sheets = read_sheets(name, data)
    fin = financials(sheets)
    return {"name": name, "sheets": sheets, "financials": fin, "verdict": verdict(fin)}


def compare(workbooks: list[dict]) -> dict:
    """Workbooks in order (oldest first). Each gets a verdict judged against the one before it;
    every ordered pair (earlier -> later) gets a cell diff and the verdict change."""
    for prev, w in zip(workbooks, workbooks[1:]):
        w["verdict"] = verdict(w["financials"], prior=prev["financials"])
    diffs = {}
    for i in range(len(workbooks)):
        for j in range(i + 1, len(workbooks)):
            a, b = workbooks[i], workbooks[j]
            diffs[f"{i}-{j}"] = {"from": i, "to": j, "cells": diff_sheets(a["sheets"], b["sheets"]),
                                 "verdict": verdict_change(a["verdict"], b["verdict"])}
    return {
        "workbooks": [{"index": i, "name": w["name"],
                       "sheets": [{"sheet": s["sheet"], "rows": s["rows"][:60]} for s in w["sheets"]],
                       "financials": {**w["financials"], "metrics": {m: METRIC_LABELS[m] for m in w["financials"]["series"]}},
                       "verdict": w["verdict"]} for i, w in enumerate(workbooks)],
        "diffs": diffs,
    }
